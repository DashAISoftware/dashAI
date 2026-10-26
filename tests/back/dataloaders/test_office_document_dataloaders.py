"""Tests for the python-docx and unstructured document dataloaders."""

import contextlib
import hashlib
import io
import os
import zipfile
from pathlib import Path

import pytest
from kink import di

from DashAI.back.dataloaders.classes import unstructured_document_dataloader as module
from DashAI.back.dataloaders.classes.document_dataloader import DocumentLoaderError
from DashAI.back.dataloaders.classes.python_docx_document_dataloader import (
    PythonDocxDocumentDataLoader,
)
from DashAI.back.dataloaders.classes.unstructured_document_dataloader import (
    PRIVACY_ENVIRONMENT,
    UnstructuredDocumentDataLoader,
)
from DashAI.back.types.dashai_image import DashAIImage
from DashAI.back.types.value_types import Text
from tests.back.dataloaders.document_samples import png_bytes, write_docx

_SENTINEL = object()


@contextlib.contextmanager
def _component_root(path):
    try:
        old = di["config"]
    except KeyError:
        old = _SENTINEL
    di["config"] = {
        **({} if old is _SENTINEL else old),
        "COMPONENT_PATH": str(path),
    }
    try:
        yield Path(path)
    finally:
        if old is _SENTINEL:
            del di["config"]
        else:
            di["config"] = old


@pytest.fixture(scope="session")
def spacy_model_root(request, tmp_path_factory):
    cache = getattr(request.config, "cache", None)
    root = (
        cache.mkdir("unstructured-spacy-model")
        if cache is not None
        else tmp_path_factory.mktemp("unstructured-spacy-model")
    )
    with _component_root(root):
        if not UnstructuredDocumentDataLoader.is_downloaded():
            try:
                UnstructuredDocumentDataLoader.download()
            except Exception as e:
                pytest.skip(f"The spaCy model could not be downloaded: {e}")
    return root


@pytest.fixture
def unstructured_loader(spacy_model_root):
    with _component_root(spacy_model_root):
        yield UnstructuredDocumentDataLoader()


@pytest.fixture
def components_root(tmp_path):
    with _component_root(tmp_path) as root:
        yield root


def _elements(document):
    return [(item.kind, item.content, item.attrs) for item in document.items()]


def _fake_wheel():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("en_core_web_sm/__init__.py", "")
        archive.writestr("en_core_web_sm-3.8.0.dist-info/METADATA", "Name: x")
    return buffer.getvalue()


class _FakeResponse:
    def __init__(self, data):
        self.data = data
        self.headers = {"content-length": str(len(data))}

    def raise_for_status(self):
        return None

    def iter_bytes(self, size):
        for start in range(0, len(self.data), size):
            yield self.data[start : start + size]


def _serve(monkeypatch, data):
    import httpx

    @contextlib.contextmanager
    def stream(method, url, **kwargs):
        assert (method, url) == ("GET", module.SPACY_MODEL_URL)
        yield _FakeResponse(data)

    monkeypatch.setattr(httpx, "stream", stream)


def test_python_docx_names_items_after_their_word_style(tmp_path):
    path = write_docx(tmp_path / "report.docx")
    document = PythonDocxDocumentDataLoader().load_document(
        path, {"extract_images": False}
    )
    assert _elements(document) == [
        ("title", "Report", {}),
        ("heading", "Introduction", {"level": 1}),
        ("paragraph", "Body text.", {}),
        ("list_item", "Bullet", {"ordered": False, "depth": 0}),
        ("list_item", "Nested bullet", {"ordered": False, "depth": 1}),
        ("list_item", "Numbered", {"ordered": True, "depth": 0}),
        ("heading", "Results", {"level": 2}),
        (
            "table",
            "Name\tScore\nAda\t",
            {"rows": [["Name", "Score"], ["Ada", ""]], "header": True},
        ),
        ("caption", "Figure 1", {}),
        ("quote", "Quoted words.", {}),
        ("paragraph", "The end.", {}),
    ]
    assert all(item.page is None for item in document.items())
    assert document.metadata["n_pages"] is None
    assert document.metadata["file_type"] == "docx"
    assert document.metadata["dataloader"] == "PythonDocxDocumentDataLoader"


def test_python_docx_puts_images_after_their_paragraph(tmp_path):
    path = write_docx(tmp_path / "report.docx")
    items = PythonDocxDocumentDataLoader().load_document(path, {}).items()
    kinds = [item.kind for item in items]
    image = items[kinds.index("image")]
    assert kinds[kinds.index("image") - 1] == "table"
    assert kinds[kinds.index("image") + 1] == "caption"
    assert image.type is DashAIImage
    assert image.attrs == {"alt": "A red square"}
    assert image.content.bytes == png_bytes()


def test_python_docx_keeps_headers_and_footers_only_when_asked(tmp_path):
    path = write_docx(tmp_path / "report.docx")
    loader = PythonDocxDocumentDataLoader()
    default = loader.load_document(path, {})
    assert "page_header" not in [item.kind for item in default.items()]
    items = loader.load_document(path, {"include_headers_footers": True}).items()
    assert (items[0].kind, items[0].content) == ("page_header", "Top")
    assert (items[-1].kind, items[-1].content) == ("page_footer", "Bottom")


def test_python_docx_wraps_parse_errors(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip")
    with pytest.raises(DocumentLoaderError, match="python-docx"):
        PythonDocxDocumentDataLoader().load_document(str(path), {})


def test_unstructured_names_items_after_their_category(tmp_path, unstructured_loader):
    path = write_docx(tmp_path / "report.docx")
    document = unstructured_loader.load_document(path, {})
    elements = _elements(document)
    assert ("heading", "Introduction", {"level": 1}) in elements
    assert ("list_item", "Bullet", {"ordered": False, "depth": 0}) in elements
    assert (
        "table",
        "Name\tScore\nAda\t",
        {"rows": [["Name", "Score"], ["Ada", ""]], "header": False},
    ) in elements
    assert all(item.type is Text for item in document.items())
    assert document.metadata["dataloader"] == "UnstructuredDocumentDataLoader"


def test_unstructured_counts_levels_from_the_shallowest_one(
    tmp_path, unstructured_loader
):
    path = tmp_path / "page.html"
    path.write_text(
        "<html><body><h1>Top</h1><h2>Inner</h2>"
        "<ul><li>one</li><li>two<ul><li>deep</li></ul></li></ul>"
        "<table><thead><tr><th>A</th><th>B</th></tr></thead>"
        "<tr><td>1</td><td>2</td></tr></table>"
        "<pre><code>x = 1</code></pre></body></html>",
        encoding="utf-8",
    )
    elements = _elements(unstructured_loader.load_document(str(path), {}))
    assert elements == [
        ("heading", "Top", {"level": 1}),
        ("heading", "Inner", {"level": 2}),
        ("list_item", "one", {"ordered": False, "depth": 0}),
        ("list_item", "two", {"ordered": False, "depth": 0}),
        ("list_item", "deep", {"ordered": False, "depth": 1}),
        ("table", "A\tB\n1\t2", {"rows": [["A", "B"], ["1", "2"]], "header": True}),
        ("code", "x = 1", {}),
    ]


def test_unstructured_uses_slides_as_pages(tmp_path, unstructured_loader):
    from pptx import Presentation

    presentation = Presentation()
    for title in ["First", "", "Third"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
    path = tmp_path / "deck.pptx"
    presentation.save(str(path))
    document = unstructured_loader.load_document(str(path), {})
    assert [(item.page, item.content) for item in document.items()] == [
        (1, "First"),
        (3, "Third"),
    ]
    assert document.metadata["n_pages"] == 3


def test_unstructured_keeps_headers_and_footers_only_when_asked(
    tmp_path, unstructured_loader
):
    path = write_docx(tmp_path / "report.docx")
    loader = unstructured_loader
    default = [item.kind for item in loader.load_document(path, {}).items()]
    assert "page_header" not in default
    assert "page_footer" not in default
    items = loader.load_document(path, {"include_headers_footers": True}).items()
    kinds = {item.kind: item.content for item in items}
    assert kinds["page_header"] == "Top"
    assert kinds["page_footer"] == "Bottom"


def test_unstructured_turns_its_analytics_off(
    tmp_path, monkeypatch, unstructured_loader
):
    for name in PRIVACY_ENVIRONMENT:
        monkeypatch.delenv(name, raising=False)
    path = tmp_path / "notes.md"
    path.write_text("# Notes\n\nText.\n", encoding="utf-8")
    unstructured_loader.load_document(str(path), {})
    for name, value in PRIVACY_ENVIRONMENT.items():
        assert os.environ[name] == value


def test_unstructured_wraps_parse_errors(tmp_path, unstructured_loader):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip")
    with pytest.raises(DocumentLoaderError, match="unstructured"):
        unstructured_loader.load_document(str(path), {})


def test_unstructured_table_rows_read_header_cells():
    rows, header = UnstructuredDocumentDataLoader.table_rows(
        "<table><tr><th>a</th><th>b</th></tr><tr><td> 1 </td><td/></tr></table>"
    )
    assert (rows, header) == ([["a", "b"], ["1", ""]], True)


def test_unstructured_refuses_to_parse_before_its_download(components_root):
    path = components_root / "notes.md"
    path.write_text("# Notes", encoding="utf-8")
    with pytest.raises(DocumentLoaderError, match="must be downloaded"):
        UnstructuredDocumentDataLoader().load_document(str(path), {})


def test_unstructured_download_unpacks_the_checked_wheel(components_root, monkeypatch):
    data = _fake_wheel()
    monkeypatch.setattr(module, "SPACY_MODEL_SHA256", hashlib.sha256(data).hexdigest())
    monkeypatch.setattr(module, "DOWNLOAD_CHUNK_BYTES", 64)
    _serve(monkeypatch, data)
    reports = []
    assert not UnstructuredDocumentDataLoader.is_downloaded()
    UnstructuredDocumentDataLoader.download(lambda *args: reports.append(args))
    target = UnstructuredDocumentDataLoader.component_dir()
    assert target == components_root / "UnstructuredDocumentDataLoader"
    assert UnstructuredDocumentDataLoader.is_downloaded()
    assert not list(target.glob("*.whl"))
    fractions = [fraction for fraction, _ in reports]
    assert fractions == sorted(set(fractions))
    assert fractions[-1] == 1.0


def test_unstructured_download_rejects_a_wheel_with_another_hash(
    components_root, monkeypatch
):
    _serve(monkeypatch, _fake_wheel())
    with pytest.raises(ValueError, match="SHA-256"):
        UnstructuredDocumentDataLoader.download()
    assert not UnstructuredDocumentDataLoader.is_downloaded()
    assert not list(UnstructuredDocumentDataLoader.component_dir().iterdir())


def test_only_unstructured_requires_a_download():
    metadata = UnstructuredDocumentDataLoader.get_metadata()
    assert metadata["requires_download"] is True
    assert metadata["download_size_bytes"] == 12806118
    assert PythonDocxDocumentDataLoader.get_metadata()["requires_download"] is False
